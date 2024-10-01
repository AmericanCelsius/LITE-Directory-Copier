import os
import threading
from queue import Queue
from docx import Document
import openpyxl
from PyPDF2 import PdfReader

# List of binary file extensions to skip
BINARY_EXTENSIONS = ['.jpg', '.jpeg', '.png', '.gif', '.bmp', '.tiff', '.mp4', '.mov', '.exe', '.zip', '.bin']

# Function to recursively list all files in the directory
def list_files_recursive(directory):
    file_list = []
    for dirpath, _, filenames in os.walk(directory):
        for filename in filenames:
            file_path = os.path.join(dirpath, filename)
            file_list.append(file_path)
    return file_list

# Function to check if a file is binary based on its extension
def is_binary_file(file_path):
    _, extension = os.path.splitext(file_path)
    return extension.lower() in BINARY_EXTENSIONS

# Function to read .docx files
def read_docx(file_path):
    try:
        doc = Document(file_path)
        return '\n'.join([para.text for para in doc.paragraphs])
    except Exception as e:
        return f"\n\nCould not read .docx file: {file_path}. Error: {e}\n\n"

# Function to read .xlsx files
def read_xlsx(file_path):
    try:
        wb = openpyxl.load_workbook(file_path)
        sheet = wb.active
        data = []
        for row in sheet.iter_rows(values_only=True):
            data.append("\t".join([str(cell) for cell in row]))
        return '\n'.join(data)
    except Exception as e:
        return f"\n\nCould not read .xlsx file: {file_path}. Error: {e}\n\n"

# Function to read .pdf files
def read_pdf(file_path):
    try:
        reader = PdfReader(file_path)
        text = ""
        for page in reader.pages:
            text += page.extract_text() + "\n"
        return text
    except Exception as e:
        return f"\n\nCould not read .pdf file: {file_path}. Error: {e}\n\n"

# Function to read file content based on the file type
def read_file(file_path):
    _, extension = os.path.splitext(file_path)

    # Check for binary files (skip)
    if is_binary_file(file_path):
        return f"\n\n--- {file_path} is a binary file and was skipped ---\n\n"

    # Handle .docx files
    if extension.lower() == '.docx':
        return read_docx(file_path)

    # Handle .xlsx files
    if extension.lower() == '.xlsx':
        return read_xlsx(file_path)

    # Handle .pdf files
    if extension.lower() == '.pdf':
        return read_pdf(file_path)

    # For text files or unknown types, read as plain text
    try:
        with open(file_path, 'r', encoding='utf-8') as infile:
            return infile.read()
    except Exception as e:
        return f"\n\nCould not read file: {file_path}. Error: {e}\n\n"

# Worker function for threading
def worker(queue, output_file_lock, output_file, total_files, progress_counter):
    while True:
        file_path = queue.get()
        if file_path is None:
            break

        # Read file content and write it to the output file
        content = read_file(file_path)
        with output_file_lock:
            with open(output_file, 'a', encoding='utf-8') as outfile:
                outfile.write(f"\n\n--- Start of {file_path} ---\n\n")
                outfile.write(content)
                outfile.write(f"\n\n--- End of {file_path} ---\n\n")

        # Update progress
        progress_counter['count'] += 1
        progress = (progress_counter['count'] / total_files) * 100
        print(f"Processing {file_path} - {progress_counter['count']} out of {total_files} files traversed ({progress:.2f}% complete)")

        queue.task_done()

# Function to collect files to a single .txt file with multithreading
def collect_files_to_txt(directory, output_file, num_threads=4):
    # Delete the existing aggregated file if it exists
    if os.path.exists(output_file):
        print(f"Deleting existing file: {output_file}")
        os.remove(output_file)

    # Start by writing the directory name and path at the top of the file
    with open(output_file, 'w', encoding='utf-8') as outfile:
        outfile.write(f"Directory Name: {os.path.basename(directory)}\n")
        outfile.write(f"Directory Path: {os.path.abspath(directory)}\n")
        outfile.write("Recursive file list:\n")

        # Recursively list all files
        file_list = list_files_recursive(directory)
        for file_path in file_list:
            outfile.write(f"{file_path}\n")

    total_files = len(file_list)  # Total number of files
    progress_counter = {'count': 0}  # To track progress across threads

    # Create a queue to process files
    queue = Queue()
    output_file_lock = threading.Lock()

    # Create worker threads
    threads = []
    for _ in range(num_threads):
        thread = threading.Thread(target=worker, args=(queue, output_file_lock, output_file, total_files, progress_counter))
        thread.start()
        threads.append(thread)

    # Add files to the queue
    for file_path in file_list:
        queue.put(file_path)

    # Wait for the queue to be processed
    queue.join()

    # Stop workers
    for _ in range(num_threads):
        queue.put(None)
    for thread in threads:
        thread.join()

# Set output file in the same directory
directory = os.path.dirname(os.path.abspath(__file__))
output_file = os.path.join(directory, "aggregated_files.txt")

# Call the function to process the files
collect_files_to_txt(directory, output_file)
